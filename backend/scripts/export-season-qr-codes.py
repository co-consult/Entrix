#!/usr/bin/env python3
"""Export season QR codes to CSV + formatted Excel (sortable tables, styled headers)."""
from __future__ import annotations

import csv
import json
import os
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

CSV_HEADERS = [
    "Plan Code",
    "Plan Name",
    "Serial Number",
    "QR Code",
    "PIN (Onboarding Key)",
    "Seat",
    "Zone",
    "Porte",
    "Status",
]

XLSX_HEADERS = [
    "Plan Code",
    "Plan Name",
    "Serial Number",
    "QR Code",
    "PIN (Onboarding Key)",
    "Seat",
    "Zone",
    "Porte",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TABLE_STYLE = TableStyleInfo(
    name="TableStyleMedium2",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)

VB_PLAN_BASES = frozenset({"GRADINS_VB", "CHAISE_VB"})


def is_vb_plan(code: str) -> bool:
    base = code.rsplit("-", 1)[0] if "-" in code else code
    return base in VB_PLAN_BASES or "_VB" in base


def fetch_rows(season: str) -> list[dict[str, str]]:
    query = f"""
    SELECT sp.code, sp.name, p.serial_number, p.qr_code, p.onboarding_key,
      COALESCE(p.metadata->>'seat', '') AS seat,
      COALESCE(p.metadata->>'zone', '') AS zone,
      COALESCE(p.metadata->>'entry_gate', '') AS porte,
      p.status::text
    FROM physical_qr_codes p
    JOIN subscription_plans sp ON sp.id = p.subscription_plan_id
    WHERE sp.metadata->>'season' = '{season}'
    ORDER BY sp.code, p.serial_number::text
    """
    raw = subprocess.check_output(
        [
            "docker",
            "exec",
            "entrix_postgres",
            "psql",
            "-U",
            "entrix_user",
            "-d",
            "entrix_db",
            "-t",
            "-A",
            "-F",
            "\t",
            "-c",
            query,
        ],
        text=True,
    )
    rows: list[dict[str, str]] = []
    for line in raw.strip().split("\n"):
        if not line.strip():
            continue
        parts = line.split("\t")
        if len(parts) < 9:
            continue
        rows.append(dict(zip(CSV_HEADERS, parts)))
    return rows


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_HEADERS)
        writer.writeheader()
        writer.writerows(rows)


def autosize_columns(ws, headers: list[str], row_count: int) -> None:
    for idx, header in enumerate(headers, start=1):
        col = get_column_letter(idx)
        max_len = len(header)
        for r in range(2, min(row_count + 2, 502)):
            val = ws.cell(row=r, column=idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col].width = min(max_len + 2, 48)


def safe_table_name(name: str, index: int) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_]", "_", name)
    if not cleaned or not cleaned[0].isalpha():
        cleaned = f"T_{cleaned}"
    return f"QR_{index}_{cleaned}"[:240]


def write_xlsx(path: Path, rows: list[dict[str, str]], sheet_name: str, table_index: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    ws.append(XLSX_HEADERS)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([row[h] for h in XLSX_HEADERS])

    last_col = get_column_letter(len(XLSX_HEADERS))
    last_row = len(rows) + 1
    table = Table(
        displayName=safe_table_name(sheet_name, table_index),
        ref=f"A1:{last_col}{last_row}",
    )
    table.tableStyleInfo = TABLE_STYLE
    ws.add_table(table)
    ws.freeze_panes = "A2"
    autosize_columns(ws, XLSX_HEADERS, len(rows))

    wb.save(path)


def main() -> int:
    season = sys.argv[1] if len(sys.argv) > 1 else "2026-2027"
    out_base = Path(
        sys.argv[2] if len(sys.argv) > 2 else f"/opt/entrix/exports/saison-{season}"
    )

    all_rows = fetch_rows(season)
    by_plan: dict[str, list[dict[str, str]]] = {}
    for row in all_rows:
        by_plan.setdefault(row["Plan Code"], []).append(row)

    manifest: dict[str, list[dict[str, str | int]]] = {"football": [], "vb": []}
    table_index = 0

    for code in sorted(by_plan.keys()):
        plan_rows = by_plan[code]
        vb = is_vb_plan(code)
        group = "vb" if vb else "football"
        out_csv = out_base / "csv" / group
        out_xlsx = out_base / "xlsx" / group
        out_csv.mkdir(parents=True, exist_ok=True)
        out_xlsx.mkdir(parents=True, exist_ok=True)

        table_index += 1
        csv_path = out_csv / f"{code}.csv"
        xlsx_path = out_xlsx / f"{code}.xlsx"
        write_csv(csv_path, plan_rows)
        write_xlsx(xlsx_path, plan_rows, code, table_index)
        manifest[group].append({
            "file": f"{group}/{code}.csv",
            "xlsx": f"{group}/{code}.xlsx",
            "plan": code,
            "count": len(plan_rows),
        })

    football_rows = [r for r in all_rows if not is_vb_plan(r["Plan Code"])]
    vb_rows = [r for r in all_rows if is_vb_plan(r["Plan Code"])]

    football_csv = out_base / "csv" / "football"
    football_xlsx = out_base / "xlsx" / "football"
    vb_csv = out_base / "csv" / "vb"
    vb_xlsx = out_base / "xlsx" / "vb"
    football_csv.mkdir(parents=True, exist_ok=True)
    football_xlsx.mkdir(parents=True, exist_ok=True)
    vb_csv.mkdir(parents=True, exist_ok=True)
    vb_xlsx.mkdir(parents=True, exist_ok=True)

    combined_football_csv = football_csv / f"_TOUS_PLANS_{season}.csv"
    combined_football_xlsx = football_xlsx / f"_TOUS_PLANS_{season}.xlsx"
    write_csv(combined_football_csv, football_rows)
    write_xlsx(combined_football_xlsx, football_rows, "Tous les plans", 0)

    combined_vb_csv = vb_csv / f"_TOUS_PLANS_VB_{season}.csv"
    combined_vb_xlsx = vb_xlsx / f"_TOUS_PLANS_VB_{season}.xlsx"
    write_csv(combined_vb_csv, vb_rows)
    write_xlsx(combined_vb_xlsx, vb_rows, "Tous les plans VB", 9000)

    manifest_path = out_base / "csv" / "_MANIFEST.json"
    with manifest_path.open("w", encoding="utf-8") as f:
        json.dump(
            {
                "season": season,
                "total": len(all_rows),
                "football": {"count": len(football_rows), "files": manifest["football"]},
                "vb": {"count": len(vb_rows), "files": manifest["vb"]},
            },
            f,
            indent=2,
        )

    print(f"Exported {len(all_rows)} rows total")
    print(f"  Football: {len(football_rows)} → {football_csv} | {football_xlsx}")
    print(f"  VB:       {len(vb_rows)} → {vb_csv} | {vb_xlsx}")

    for folder in (out_base / "csv", out_base / "xlsx"):
        for stale in list(folder.glob("*.csv")) + list(folder.glob("*.xlsx")):
            stale.unlink()

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
